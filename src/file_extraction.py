from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from PIL import Image
import fitz
import pytesseract
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError


SUPPORTED_FILE_TYPES = [
    "txt",
    "pdf",
    "docx",
    "png",
    "jpg",
    "jpeg",
    "bmp",
    "tif",
    "tiff",
]


class OcrError(RuntimeError):
    """Raised when OCR cannot be performed for an uploaded file."""


def _ocr_image(image):
    try:
        return pytesseract.image_to_string(image).strip()
    except pytesseract.TesseractNotFoundError as error:
        raise OcrError(
            "OCR requires Tesseract to be installed and available on the "
            "system PATH."
        ) from error
    except pytesseract.TesseractError as error:
        raise OcrError(f"OCR failed: {error}") from error


def _ocr_pdf(file_bytes, page_indexes=None):
    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
        page_text = []
        pages_to_read = (
            enumerate(document)
            if page_indexes is None
            else (
                (index, document[index])
                for index in page_indexes
            )
        )
        for _, page in pages_to_read:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes(
                "RGB",
                (pixmap.width, pixmap.height),
                pixmap.samples,
            )
            text = _ocr_image(image)
            if text:
                page_text.append(text)
        return "\n\n".join(page_text)
    except OcrError:
        raise
    except (fitz.FileDataError, RuntimeError) as error:
        raise OcrError(f"Could not render this PDF for OCR: {error}") from error


def _extract_docx_text(document):
    blocks = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if cells:
                blocks.append(" | ".join(cells))
    for shape in document.inline_shapes:
        relationship_id = (
            shape._inline.graphic.graphicData.pic.pic.blipFill.blip.embed
        )
        image_part = document.part.related_parts[relationship_id]
        image = Image.open(BytesIO(image_part.blob))
        text = _ocr_image(image)
        if text:
            blocks.append(text)
    return "\n\n".join(blocks)


def extract_text(uploaded_file) -> str:
    """Extract plain text from a supported Streamlit uploaded file."""
    extension = Path(uploaded_file.name).suffix.lower()
    file_bytes = uploaded_file.getvalue()

    if extension == ".txt":
        text = file_bytes.decode("utf-8-sig")
        if not text.strip():
            raise ValueError("This TXT file contains no readable text.")
        return text

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        if reader.is_encrypted:
            raise FileNotDecryptedError(
                "This PDF is password-protected and cannot be read. "
                "Remove the password or export an unprotected copy."
            )
        pages = [page.extract_text() or "" for page in reader.pages]
        text_pages = [page.strip() for page in pages if page.strip()]
        missing_page_indexes = [
            index for index, page in enumerate(pages) if not page.strip()
        ]
        if missing_page_indexes:
            ocr_text = _ocr_pdf(file_bytes, missing_page_indexes)
            if ocr_text:
                text_pages.append(ocr_text)
        text = "\n\n".join(text_pages)
        if not text.strip():
            raise ValueError("This PDF contains no readable text.")
        return text

    if extension == ".docx":
        try:
            document = Document(BytesIO(file_bytes))
        except BadZipFile as error:
            raise ValueError(
                "This DOCX file is corrupt or not a valid Word document."
            ) from error

        text = _extract_docx_text(document)
        if not text.strip():
            raise ValueError(
                "This DOCX file contains no readable text."
            )
        return text

    if extension in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        try:
            text = _ocr_image(Image.open(BytesIO(file_bytes)))
        except (OSError, ValueError) as error:
            raise ValueError("This image file is corrupt or unreadable.") from error
        if not text:
            raise ValueError("This image contains no readable text.")
        return text

    raise ValueError(
        f"Unsupported file type '{extension or 'unknown'}'. "
        "Please upload a TXT, PDF, DOCX, or image file."
    )
